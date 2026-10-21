"""Shared fixtures: sample PDFs and images made on the fly (no binary test
files in the repo). Qt runs offscreen, so the tests need no display."""

from __future__ import annotations

import os
import sys
from pathlib import Path

import pytest

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")


@pytest.fixture(scope="session")
def qapp():
    from PyQt6.QtGui import QGuiApplication

    app = QGuiApplication.instance() or QGuiApplication([])
    yield app


def make_pdf(path: Path, pages: int = 3, password: str | None = None) -> Path:
    import pymupdf

    doc = pymupdf.open()
    for i in range(pages):
        page = doc.new_page(width=595, height=842)  # A4
        page.insert_text((72, 100), f"Hello page {i + 1}", fontsize=24)
        page.insert_text((72, 140), "Spindle test document", fontsize=12)
    if password:
        doc.save(path, encryption=pymupdf.PDF_ENCRYPT_AES_256, user_pw=password, owner_pw=password)
    else:
        doc.save(path)
    doc.close()
    return path


def make_image(path: Path, fmt: str, size=(64, 48), alpha: bool = False) -> Path:
    from PyQt6.QtGui import QColor, QImage

    image = QImage(size[0], size[1], QImage.Format.Format_ARGB32 if alpha else QImage.Format.Format_RGB32)
    image.fill(QColor(0, 0, 0, 0) if alpha else QColor("#3a7bd5"))
    if alpha:  # opaque square in the middle, transparent border
        for y in range(size[1] // 4, 3 * size[1] // 4):
            for x in range(size[0] // 4, 3 * size[0] // 4):
                image.setPixelColor(x, y, QColor("#e94e1b"))
    assert image.save(str(path), fmt), f"could not write {fmt}"
    return path


def add_exif_orientation(jpeg: Path, orientation: int) -> None:
    """Insert a minimal EXIF segment with the Orientation tag into a JPEG."""
    tiff = (b"II*\x00" + (8).to_bytes(4, "little")           # header, IFD at 8
            + (1).to_bytes(2, "little")                       # 1 entry
            + (0x0112).to_bytes(2, "little") + (3).to_bytes(2, "little")
            + (1).to_bytes(4, "little") + orientation.to_bytes(2, "little") + b"\x00\x00"
            + (0).to_bytes(4, "little"))                      # no next IFD
    payload = b"Exif\x00\x00" + tiff
    segment = b"\xff\xe1" + (len(payload) + 2).to_bytes(2, "big") + payload
    data = jpeg.read_bytes()
    assert data[:2] == b"\xff\xd8"
    jpeg.write_bytes(data[:2] + segment + data[2:])


def make_photo(path: Path, size=(1600, 1200), fmt: str = "JPG", quality: int = 95) -> Path:
    """A noisy, photo-like image (compresses like a real photo)."""
    import random

    from PyQt6.QtGui import QColor, QImage, QImageWriter, QLinearGradient, QPainter

    rnd = random.Random(7)
    image = QImage(size[0], size[1], QImage.Format.Format_RGB32)
    p = QPainter(image)
    g = QLinearGradient(0, 0, size[0], size[1])
    g.setColorAt(0, QColor("#2b5fa8"))
    g.setColorAt(1, QColor("#f5a25d"))
    p.fillRect(image.rect(), g)
    for _ in range(4000):
        p.fillRect(rnd.randrange(size[0]), rnd.randrange(size[1]), rnd.randrange(2, 12), rnd.randrange(2, 12),
                   QColor(rnd.randrange(256), rnd.randrange(256), rnd.randrange(256)))
    p.end()
    writer = QImageWriter(str(path), fmt.encode())
    writer.setQuality(quality)
    assert writer.write(image)
    return path


def make_photo_pdf(path: Path, photo: Path) -> Path:
    """A one-page PDF with a big JPEG photo shown small (≈ 580 DPI)."""
    import pymupdf

    doc = pymupdf.open()
    page = doc.new_page()
    page.insert_text((72, 60), "Photo report", fontsize=20)
    page.insert_image(pymupdf.Rect(72, 80, 272, 230), filename=str(photo))
    doc.save(path, garbage=3, deflate=True)
    doc.close()
    return path


def jpeg_segments(data: bytes) -> list[tuple[int, bytes]]:
    """(marker, payload) of every segment before the image data."""
    import struct

    out, pos = [], 2
    while pos < len(data) and data[pos] == 0xFF:
        marker = data[pos + 1]
        if marker == 0xDA:
            break
        length = struct.unpack(">H", data[pos + 2:pos + 4])[0]
        out.append((marker, data[pos + 4:pos + 2 + length]))
        pos += 2 + length
    return out


def jpeg_scan(data: bytes) -> bytes:
    """Everything from the start-of-scan marker on: the actual image data."""
    return data[data.index(b"\xff\xda"):]


def add_jpeg_segment(jpeg: Path, marker: int, payload: bytes) -> None:
    import struct

    data = jpeg.read_bytes()
    seg = bytes((0xFF, marker)) + struct.pack(">H", len(payload) + 2) + payload
    jpeg.write_bytes(data[:2] + seg + data[2:])


# ---- multi-frame images (built by hand: Qt can't write any of these) ----

def make_multipage_tiff(path: Path, pages: int = 2, size=(4, 3)) -> Path:
    """An uncompressed 8-bit grayscale TIFF with `pages` pages."""
    import struct

    w, h = size
    ifd_size = 2 + 12 * 9 + 4
    data_start = 8 + ifd_size * pages
    out = bytearray(b"II*\x00" + struct.pack("<I", 8))
    for p in range(pages):
        entries = [(256, 3, w), (257, 3, h), (258, 3, 8), (259, 3, 1), (262, 3, 1),
                   (273, 4, data_start + p * w * h), (277, 3, 1), (278, 3, h), (279, 4, w * h)]
        out += struct.pack("<H", len(entries))
        for tag, typ, val in entries:
            out += struct.pack("<HHI", tag, typ, 1)
            out += struct.pack("<HH", val, 0) if typ == 3 else struct.pack("<I", val)
        out += struct.pack("<I", 8 + ifd_size * (p + 1) if p < pages - 1 else 0)
    for p in range(pages):
        out += bytes([40 + 60 * p]) * (w * h)
    path.write_bytes(bytes(out))
    return path


def make_animated_gif(path: Path, frames: int = 2) -> Path:
    frame = b"!\xf9\x04\x01\x0a\x00\x00\x00,\x00\x00\x00\x00\x01\x00\x01\x00\x00\x02\x02D\x01\x00"
    path.write_bytes(b"GIF89a\x01\x00\x01\x00\x80\x00\x00\xff\xff\xff\x00\x00\x00"
                     b"!\xff\x0bNETSCAPE2.0\x03\x01\x00\x00\x00" + frame * frames + b";")
    return path


def make_animated_webp(path: Path, frames: int = 2, size=(16, 16)) -> Path:
    """Animated WEBP: VP8X + ANIM + one ANMF per frame, reusing a still
    frame written by Qt."""
    import struct

    from PyQt6.QtGui import QColor, QImage

    still = path.with_name("still-frame.webp")
    img = QImage(size[0], size[1], QImage.Format.Format_ARGB32)
    img.fill(QColor("#3a7bd5"))
    assert img.save(str(still), "WEBP")
    raw = still.read_bytes()
    still.unlink()
    frame_chunks = raw[12:]  # VP8 / VP8L chunk(s) of the still image

    def chunk(ctype: bytes, payload: bytes) -> bytes:
        return ctype + struct.pack("<I", len(payload)) + payload + (b"\x00" if len(payload) & 1 else b"")

    w3 = (size[0] - 1).to_bytes(3, "little")
    h3 = (size[1] - 1).to_bytes(3, "little")
    body = chunk(b"VP8X", bytes([0x12, 0, 0, 0]) + w3 + h3)  # animation + alpha flags
    body += chunk(b"ANIM", b"\x00\x00\x00\x00" + b"\x00\x00")
    for _ in range(frames):
        body += chunk(b"ANMF", b"\x00\x00\x00" * 2 + w3 + h3 + (100).to_bytes(3, "little") + b"\x00" + frame_chunks)
    path.write_bytes(b"RIFF" + struct.pack("<I", len(body) + 4) + b"WEBP" + body)
    return path


def make_apng(path: Path) -> Path:
    """A PNG marked as animated (acTL chunk before the image data)."""
    import struct
    import zlib

    make_image(path, "PNG", alpha=True)
    data = path.read_bytes()
    payload = struct.pack(">II", 2, 0)
    actl = struct.pack(">I", 8) + b"acTL" + payload + struct.pack(">I", zlib.crc32(b"acTL" + payload))
    at = data.index(b"IDAT") - 4
    path.write_bytes(data[:at] + actl + data[at:])
    return path


def append_motion_video(jpeg: Path, secret: bytes = b"GPS 12.9716N 77.5946E") -> None:
    """Append a fake MP4 the way Google motion photos do (after the JPEG)."""
    import struct

    ftyp = struct.pack(">I", 24) + b"ftypmp42" + b"\x00\x00\x00\x00mp42isom"
    jpeg.write_bytes(jpeg.read_bytes() + ftyp + struct.pack(">I", 8 + len(secret)) + b"mdat" + secret)


def add_exif_gps(jpeg: Path, text: str) -> None:
    """Insert an EXIF block with a GPS sub-block (latitude ref + a map
    datum string holding `text`), like a phone camera writes."""
    import struct

    data = text.encode() + b"\x00"
    # IFD0 at 8: one entry, GPSInfo -> GPS IFD at 26
    ifd0 = struct.pack("<H", 1) + struct.pack("<HHII", 0x8825, 4, 1, 26) + struct.pack("<I", 0)
    gps_ifd_at = 8 + len(ifd0)
    text_at = gps_ifd_at + 2 + 2 * 12 + 4
    gps = (struct.pack("<H", 2)
           + struct.pack("<HHI", 0x0001, 2, 2) + b"N\x00\x00\x00"           # GPSLatitudeRef
           + struct.pack("<HHII", 0x0012, 2, len(data), text_at)            # GPSMapDatum
           + struct.pack("<I", 0))
    tiff = b"II*\x00" + struct.pack("<I", 8) + ifd0 + gps + data
    assert gps_ifd_at == 26
    payload = b"Exif\x00\x00" + tiff
    segment = b"\xff\xe1" + struct.pack(">H", len(payload) + 2) + payload
    raw = jpeg.read_bytes()
    jpeg.write_bytes(raw[:2] + segment + raw[2:])


def camera_exif(thumbnail: bytes | None = None, artist: str = "Ravi Tej", orientation: int = 1,
                size: tuple[int, int] | None = None) -> bytes:
    """EXIF like a phone camera writes (TIFF data, no "Exif" prefix):
    camera, dates, exposure, maker notes, GPS (Bengaluru), and optionally
    a thumbnail. Written by Pillow (independent of the editor's writer);
    the thumbnail is added with the editor's EXIF module (Pillow can't)."""
    from PIL import ExifTags, Image
    from PIL.TiffImagePlugin import IFDRational as R

    ex = Image.Exif()
    ex[0x010F] = "Google"
    ex[0x0110] = "Pixel 8"
    ex[0x013B] = artist.encode()  # as UTF-8 bytes: Pillow would write str as ASCII ("Zo?")
    ex[0x0131] = "HDR+ 1.0"
    ex[0x0112] = orientation
    e = ex.get_ifd(ExifTags.IFD.Exif)
    if size:
        e[0xA002], e[0xA003] = size
    e[0x9003] = "2024:05:01 10:00:00"
    e[0x829A] = R(1, 120)
    e[0x829D] = R(18, 10)
    e[0x927C] = b"MAKERNOTE serial 12345"
    g = ex.get_ifd(ExifTags.IFD.GPSInfo)
    g[1] = "N"
    g[2] = (12.0, 58.0, 17.76)
    g[3] = "E"
    g[4] = (77.0, 35.0, 40.56)
    g[6] = 920.5
    tiff = ex.tobytes()[6:]
    if thumbnail:
        from app.editor.exif import Exif

        parsed = Exif.parse(tiff)
        parsed.thumbnail = thumbnail
        parsed.set_short("1", 0x0103, 6)  # compression: JPEG
        tiff = parsed.to_bytes()
    return tiff


def make_camera_photo(path: Path, size=(320, 240), thumbnail: bool = True, orientation: int = 1) -> Path:
    """A JPEG with full camera EXIF (GPS, thumbnail), XMP, a comment."""
    make_photo(path, size=size)
    thumb = None
    if thumbnail:
        small = path.with_name("thumb-src.jpg")
        make_image(small, "JPG", size=(80, 60))
        thumb = small.read_bytes()
        small.unlink()
    add_jpeg_segment(path, 0xE1, b"http://ns.adobe.com/xap/1.0/\x00<x:xmpmeta xmlns:x=\"adobe:ns:meta/\">"
                     b"<rdf:RDF><rdf:Description tiff:Orientation=\"%d\" exif:GPSLatitude=\"12,58.296N\"" % orientation
                     + b" exif:GPSLongitude=\"77,35.676E\"><exif:GPSAltitude>9205/10</exif:GPSAltitude>"
                     b"<dc:title>Trip</dc:title></rdf:Description></rdf:RDF></x:xmpmeta>")
    payload = b"Exif\x00\x00" + camera_exif(thumb, orientation=orientation, size=size)
    add_jpeg_segment(path, 0xE1, payload)  # first, like cameras write it (Qt reads only the first APP1)
    add_jpeg_segment(path, 0xFE, b"taken at home")  # (comment before EXIF: still read fine)
    return path


def make_rich_pdf(path: Path, photo: Path) -> Path:
    """A PDF full of metadata: document info, XMP, a phone photo with its
    EXIF (GPS) left in, a comment with an author, an attached file, a page
    thumbnail-free page and private app data."""
    import pymupdf

    doc = pymupdf.open()
    page = doc.new_page()
    page.insert_text((72, 60), "Site visit report", fontsize=20)
    page.insert_image(pymupdf.Rect(72, 80, 272, 230), stream=photo.read_bytes())  # bytes as they are
    note = page.add_text_annot((300, 100), "Check this")
    note.set_info(title="Dr Who")
    note.update()
    doc.embfile_add("notes.txt", b"private notes", filename="notes.txt")
    doc.set_metadata({"author": "Dr Who", "title": "Site visit", "creator": "Word", "producer": "macOS Quartz",
                      "creationDate": "D:20240501100000"})
    doc.set_xml_metadata('<x:xmpmeta xmlns:x="adobe:ns:meta/"><dc:creator>Dr Who</dc:creator></x:xmpmeta>')
    doc.xref_set_key(page.xref, "PieceInfo", "<</Illustrator <</Private (secret working file)>>>>")
    info = int(doc.xref_get_key(-1, "Info")[1].split()[0])
    doc.xref_set_key(info, "Company", "(Acme Secret Corp)")  # Word adds entries like this
    doc.save(path)
    doc.close()
    return path


def make_heic(path: Path, size=(640, 480), orientation: int = 1, gps: bool = True, images: int = 1,
              icc: bytes | None = None) -> Path:
    """An iPhone-like HEIC: left half green, right half red, camera EXIF
    (GPS) and XMP. Written with pillow-heif (a test-only dependency: the
    app itself only decodes HEIC)."""
    import pillow_heif
    from PIL import Image

    im = Image.new("RGB", size, (220, 40, 40))
    im.paste((30, 190, 60), (0, 0, size[0] // 2, size[1]))
    exif = camera_exif(orientation=orientation) if gps else None
    heif = pillow_heif.from_pillow(im)
    for _ in range(images - 1):
        heif.add_from_pillow(im)
    kwargs = {"quality": 90}
    if exif:
        kwargs["exif"] = b"Exif\x00\x00" + exif
    kwargs["xmp"] = b'<x:xmpmeta xmlns:x="adobe:ns:meta/"><rdf:RDF><rdf:Description exif:GPSLatitude="12,58.296N"/></rdf:RDF></x:xmpmeta>'
    if icc:
        heif[0].info["icc_profile"] = icc
    heif.save(str(path), **kwargs)
    return path


# ---- office documents (made by hand: no office suite needed to write them) ----

def make_docx(path: Path, text: str = "Hello Spindle DOCX") -> Path:
    import docx  # python-docx (installed with pdf2docx)

    d = docx.Document()
    d.add_heading("Spindle test", level=1)
    d.add_paragraph(text)
    d.save(str(path))
    return path


_ODF = {
    ".odt": ("text", '<office:text><text:p>{text}</text:p></office:text>'),
    ".ods": ("spreadsheet", '<office:spreadsheet><table:table table:name="Sheet1"><table:table-row>'
                            '<table:table-cell office:value-type="string"><text:p>{text}</text:p>'
                            '</table:table-cell></table:table-row></table:table></office:spreadsheet>'),
    ".odp": ("presentation", '<office:presentation><draw:page draw:name="page1" draw:master-page-name="Default">'
                             '<draw:frame svg:x="2cm" '
                             'svg:y="2cm" svg:width="20cm" svg:height="3cm"><draw:text-box><text:p>{text}</text:p>'
                             '</draw:text-box></draw:frame></draw:page></office:presentation>'),
}


def make_odf(path: Path, text: str = "Hello Spindle ODF", encrypted: bool = False) -> Path:
    import zipfile

    kind, body = _ODF[path.suffix]
    ns = ('xmlns:office="urn:oasis:names:tc:opendocument:xmlns:office:1.0" '
          'xmlns:text="urn:oasis:names:tc:opendocument:xmlns:text:1.0" '
          'xmlns:table="urn:oasis:names:tc:opendocument:xmlns:table:1.0" '
          'xmlns:draw="urn:oasis:names:tc:opendocument:xmlns:drawing:1.0" '
          'xmlns:svg="urn:oasis:names:tc:opendocument:xmlns:svg-compatible:1.0" '
          'xmlns:presentation="urn:oasis:names:tc:opendocument:xmlns:presentation:1.0"')
    content = (f'<?xml version="1.0" encoding="UTF-8"?><office:document-content {ns} office:version="1.2">'
               f'<office:body>{body.format(text=text)}</office:body></office:document-content>')
    mime = f"application/vnd.oasis.opendocument.{kind}"
    crypt = ('<manifest:encryption-data manifest:checksum-type="SHA1/1K" manifest:checksum="x"/>'
             if encrypted else "")
    manifest = ('<?xml version="1.0" encoding="UTF-8"?><manifest:manifest '
                'xmlns:manifest="urn:oasis:names:tc:opendocument:xmlns:manifest:1.0" manifest:version="1.2">'
                f'<manifest:file-entry manifest:full-path="/" manifest:media-type="{mime}"/>'
                f'<manifest:file-entry manifest:full-path="content.xml" manifest:media-type="text/xml">{crypt}'
                '</manifest:file-entry></manifest:manifest>')
    # Styles with a page layout and a master page: older LibreOffice (7.3)
    # won't load a presentation without them.
    styles = (f'<?xml version="1.0" encoding="UTF-8"?><office:document-styles {ns} '
              'xmlns:style="urn:oasis:names:tc:opendocument:xmlns:style:1.0" '
              'xmlns:fo="urn:oasis:names:tc:opendocument:xmlns:xsl-fo-compatible:1.0" office:version="1.2">'
              '<office:automatic-styles><style:page-layout style:name="PM1"><style:page-layout-properties '
              'fo:page-width="28cm" fo:page-height="21cm" fo:margin-top="1cm" fo:margin-bottom="1cm" '
              'fo:margin-left="1cm" fo:margin-right="1cm"/></style:page-layout></office:automatic-styles>'
              '<office:master-styles><style:master-page style:name="Default" style:page-layout-name="PM1"/>'
              '</office:master-styles></office:document-styles>')
    manifest = manifest.replace('</manifest:manifest>', '<manifest:file-entry manifest:full-path="styles.xml" '
                                'manifest:media-type="text/xml"/></manifest:manifest>')
    with zipfile.ZipFile(path, "w") as z:
        z.writestr(zipfile.ZipInfo("mimetype"), mime)  # first, stored
        z.writestr("META-INF/manifest.xml", manifest, compress_type=zipfile.ZIP_DEFLATED)
        z.writestr("content.xml", content, compress_type=zipfile.ZIP_DEFLATED)
        z.writestr("styles.xml", styles, compress_type=zipfile.ZIP_DEFLATED)
    return path


def make_xlsx(path: Path, text: str = "Hello Spindle XLSX") -> Path:
    import zipfile

    parts = {
        "[Content_Types].xml": '<?xml version="1.0" encoding="UTF-8"?><Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types"><Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/><Default Extension="xml" ContentType="application/xml"/><Override PartName="/xl/workbook.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet.main+xml"/><Override PartName="/xl/worksheets/sheet1.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.worksheet+xml"/></Types>',
        "_rels/.rels": '<?xml version="1.0" encoding="UTF-8"?><Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships"><Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument" Target="xl/workbook.xml"/></Relationships>',
        "xl/workbook.xml": '<?xml version="1.0" encoding="UTF-8"?><workbook xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main" xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships"><sheets><sheet name="Sheet1" sheetId="1" r:id="rId1"/></sheets></workbook>',
        "xl/_rels/workbook.xml.rels": '<?xml version="1.0" encoding="UTF-8"?><Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships"><Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/worksheet" Target="worksheets/sheet1.xml"/></Relationships>',
        "xl/worksheets/sheet1.xml": f'<?xml version="1.0" encoding="UTF-8"?><worksheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main"><sheetData><row r="1"><c r="A1" t="inlineStr"><is><t>{text}</t></is></c></row></sheetData></worksheet>',
    }
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as z:
        for name, data in parts.items():
            z.writestr(name, data)
    return path
